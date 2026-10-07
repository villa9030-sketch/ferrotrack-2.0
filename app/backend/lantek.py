"""Lantek: lettura (SOLA LETTURA) del suo database e file da importare in Lantek.

Diagnostica del 06-07/10/2026 (memoria "lantek-database"): Lantek tiene i dati
in SQL Server Express sul PC server, istanza localhost\\LANTEK, database LSDB.
FerroTrack non scrive MAI in Lantek: legge l'archivio pezzi e prepara i file
che Stefano importa con i programmi di Lantek:

- i DXF con le scritte QTA/MAT/SP/ORD (importatore DXF del MES): pezzi con
  materiale e spessore giusti;
- l'XML degli ordini di produzione (XmlImporter): quantita', ordine, cliente,
  consegna. Prima si scrivevano a mano in "Invia a produzione" (ordine 1184:
  57 codici).

(L'Excel per iErp resta solo come endpoint: il modulo d'import di iErp non e'
nella licenza.)

Se Lantek non risponde (spento, aggiornato, configurazione diversa) i file si
preparano lo stesso con i dati di FerroTrack, senza i controlli: FerroTrack
funziona come prima.
"""
from __future__ import annotations

import logging
import re
import threading
import time
from datetime import date, datetime

logger = logging.getLogger(__name__)

# Collegamento predefinito; si cambia in app_config.json, chiave "lantek_db":
# {"abilitato": true, "server": "localhost\\LANTEK", "database": "LSDB"}
DB_DEFAULT = {'abilitato': True, 'server': r'localhost\LANTEK', 'database': 'LSDB',
              'driver': 'ODBC Driver 17 for SQL Server'}

# Materiali di FerroTrack -> nomi dei materiali in Lantek (DIS_MMTT_MMTT_00000100).
# Vale solo per i codici NUOVI: per quelli gia' in Lantek si usa il suo.
MATERIALI_LANTEK = {
    'S235': 'FERRO', 'S275': 'FERRO', 'S355': 'FERRO', 'FERRO': 'FERRO', 'DC01': 'FERRO',
    'C75': 'FERRO',
    'INOX_304': 'INOX', 'INOX_316': 'INOX', 'INOX': 'INOX',
    'ALU': 'ALLUMINIO', 'ALLUMINIO': 'ALLUMINIO',
    'ZINCATO': 'ZINCATO',
}

INTESTAZIONE = ['Codice Articolo', 'Quantità', 'Materiale', 'Spessore', 'Data consegna', 'Commessa']

_RX_REV = re.compile(r'^(?P<base>.+)-(?P<rev>\d{2})$')


# ---------------------------------------------------------------------------
# Lettura del database di Lantek
# ---------------------------------------------------------------------------
def _config() -> dict:
    try:
        from .database import ConfigManager
        cfg = (ConfigManager.load_config() or {}).get('lantek_db') or {}
    except Exception:
        cfg = {}
    return {**DB_DEFAULT, **(cfg if isinstance(cfg, dict) else {})}


def _connessione(cfg: dict):
    import pyodbc
    cs = (f"DRIVER={{{cfg['driver']}}};SERVER={cfg['server']};DATABASE={cfg['database']};"
          'Trusted_Connection=yes;ApplicationIntent=ReadOnly;APP=FerroTrack-lettura')
    con = pyodbc.connect(cs, timeout=4, readonly=True, autocommit=True)
    cur = con.cursor()
    # mai bloccare Lantek: si legge senza aspettare chi sta scrivendo
    cur.execute('SET TRANSACTION ISOLATION LEVEL READ UNCOMMITTED')
    return con, cur


_CACHE = {}
_LOCK = threading.Lock()
_DURATA_CACHE_S = 60


def pezzi_in_lantek(codici: list) -> dict:
    """Per ogni codice: {'esiste', 'materiale', 'spessore', 'revisioni': [...]}
    letti dall'archivio pezzi di Lantek (PPRR_PPRR_00000100), piu' le altre
    revisioni dello stesso codice (es. -00 e -02).

    Ritorna {'disponibile': bool, 'errore': str|None, 'pezzi': {...}}.
    Non solleva mai: senza Lantek, disponibile=False."""
    codici = sorted({str(c).strip() for c in codici if c and str(c).strip()})
    cfg = _config()
    if not cfg.get('abilitato'):
        return {'disponibile': False, 'errore': 'collegamento a Lantek disattivato', 'pezzi': {}}
    if not codici:
        return {'disponibile': True, 'errore': None, 'pezzi': {}}
    chiave = (cfg['server'], cfg['database'], tuple(codici))
    with _LOCK:
        c = _CACHE.get(chiave)
        if c and time.time() - c[0] < _DURATA_CACHE_S:
            return c[1]
    basi = {}
    for cod in codici:
        m = _RX_REV.match(cod)
        basi[cod] = m.group('base') if m else cod
    try:
        con, cur = _connessione(cfg)
        try:
            revisioni = {}
            for base in sorted(set(basi.values())):
                cur.execute('SELECT PrdRef FROM PPRR_PPRR_00000100 WHERE PrdRef LIKE ?',
                            [base.replace('[', '[[]') + '-[0-9][0-9]'])
                revisioni[base] = sorted(str(r[0]).strip() for r in cur.fetchall())
            # i codici chiesti e tutte le loro revisioni, a blocchi (SQL Server
            # accetta al massimo ~2100 parametri per richiesta)
            esatti = sorted(set(codici) | {r for v in revisioni.values() for r in v})
            trovati = {}
            for i in range(0, len(esatti), 500):
                blocco = esatti[i:i + 500]
                cur.execute('SELECT PrdRef, DIS_MatRef, DIS_Thickness FROM PPRR_PPRR_00000100 '
                            'WHERE PrdRef IN (' + ','.join('?' * len(blocco)) + ')', blocco)
                for prd, mat, sp in cur.fetchall():
                    trovati[str(prd).strip().upper()] = (str(prd).strip(), mat, sp)
        finally:
            con.close()
    except Exception as e:
        logger.info('Lantek non leggibile: %s', e)
        return {'disponibile': False, 'errore': 'Lantek non risponde: ' + str(e)[:160], 'pezzi': {}}
    pezzi = {}
    for cod in codici:
        t = trovati.get(cod.upper())
        altre = [r for r in revisioni.get(basi[cod], []) if r.upper() != cod.upper()]
        # Codice scritto senza revisione (25NDSPA1979) e in Lantek ce n'e' una
        # sola (25NDSPA1979-00): e' quella. Con piu' revisioni non si sceglie.
        if t is None and not _RX_REV.match(cod) and len(altre) == 1:
            t = trovati.get(altre[0].upper())
        pezzi[cod] = {
            'esiste': t is not None,
            'codice_lantek': t[0] if t else None,
            'materiale': (t[1] or None) if t else None,
            'spessore': float(t[2]) if t and t[2] is not None else None,
            'revisioni': altre,
        }
    esito = {'disponibile': True, 'errore': None, 'pezzi': pezzi}
    with _LOCK:
        if len(_CACHE) > 200:
            _CACHE.clear()
        _CACHE[chiave] = (time.time(), esito)
    return esito


# ---------------------------------------------------------------------------
# Righe del file delle quantita'
# ---------------------------------------------------------------------------
def _rev(codice: str):
    m = _RX_REV.match(codice or '')
    return (m.group('base'), int(m.group('rev'))) if m else (codice, None)


def righe_quantita(distinta: list, lantek: dict | None = None) -> list:
    """Dalle righe della distinta (app._distinta_ordine) le righe da importare:
    solo le lamiere col codice, una per codice con le quantita' sommate.

    Con i dati di Lantek (pezzi_in_lantek): materiale e spessore come sono in
    Lantek per i codici che ci sono gia', e gli avvisi (codice nuovo,
    materiale/spessore diversi, altra revisione)."""
    per_codice = {}
    for r in distinta or []:
        if r.get('tipo') != 'lamiera':
            continue
        cod = str(r.get('codice') or '').strip()
        if not cod:
            continue
        q = int(r.get('quantita') or 0)
        if q <= 0:
            continue
        x = per_codice.setdefault(cod, {'codice': cod, 'quantita': 0, 'materiale_ft': r.get('materiale'),
                                        'spessore_ft': r.get('spessore_mm')})
        x['quantita'] += q
    pezzi = (lantek or {}).get('pezzi') or {}
    disponibile = bool((lantek or {}).get('disponibile'))
    out = []
    for cod in sorted(per_codice, key=str.lower):
        x = per_codice[cod]
        mat_ft = (x['materiale_ft'] or '').strip().upper()
        mat_lt = MATERIALI_LANTEK.get(mat_ft)
        sp_ft = x['spessore_ft']
        try:
            sp_ft = float(sp_ft) if sp_ft not in (None, '') else None
        except (TypeError, ValueError):
            sp_ft = None
        avvisi = []
        p = pezzi.get(cod) if disponibile else None
        materiale, spessore = mat_lt or (x['materiale_ft'] or ''), sp_ft
        if p and p['esiste']:
            stato = 'in_lantek'
            if p['materiale']:
                if mat_lt and p['materiale'].upper() != mat_lt.upper() and not p['materiale'].upper().endswith(mat_lt.upper()):
                    avvisi.append(f"materiale: l'ordine dice {x['materiale_ft']}, in Lantek e' {p['materiale']}")
                materiale = p['materiale']
            if p['spessore'] is not None:
                if sp_ft is not None and abs(p['spessore'] - sp_ft) > 0.01:
                    avvisi.append(f"spessore: l'ordine dice {sp_ft:g} mm, in Lantek e' {p['spessore']:g} mm")
                spessore = p['spessore']
        elif disponibile:
            stato = 'nuovo'
            if p and p['revisioni'] and _rev(cod)[1] is None:
                # codice senza revisione e in Lantek ce ne sono piu' d'una
                avvisi.append('in Lantek ci sono piu\' revisioni (' + ', '.join(p['revisioni'])
                              + '): scegli quella giusta')
            else:
                avvisi.append('non ancora in Lantek: importa prima il disegno')
        else:
            stato = 'sconosciuto'
        # nel file va il codice esatto di Lantek (25NDSPA1979 -> 25NDSPA1979-00)
        codice_lt = (p.get('codice_lantek') if p else None) or cod
        if p and p['revisioni'] and _rev(cod)[1] is not None:
            rev = _rev(cod)[1]
            piu_nuove = [r for r in p['revisioni'] if _rev(r)[1] is not None and _rev(r)[1] > rev]
            if piu_nuove:
                avvisi.append('in Lantek c\'e\' una revisione piu\' nuova: ' + ', '.join(piu_nuove))
            elif stato == 'nuovo':
                avvisi.append('in Lantek ci sono altre revisioni: ' + ', '.join(p['revisioni']))
        if not mat_lt and not (p and p['esiste'] and p['materiale']):
            avvisi.append(f"materiale {x['materiale_ft'] or '?'}: da scegliere in Lantek")
        out.append({'codice': codice_lt, 'codice_ft': cod, 'quantita': x['quantita'],
                    'materiale': materiale, 'spessore': spessore, 'stato': stato, 'avvisi': avvisi})
    return out


def commessa_lantek(numero: str) -> str:
    """Il numero d'ordine come lo si scrive in Lantek (campo "riferimento"):
    "A 001252" -> "1252" (DECA: sigla, spazio, zeri). Senza sigla resta com'e'
    ("1184", "061": in Lantek lo zero c'e'); le sigle interne
    (PREV-2026-0005) restano come sono."""
    n = str(numero or '').strip()
    m = re.fullmatch(r'[A-Za-z]{1,3}[\s.\-/]+0*(\d+)', n)
    return m.group(1) if m else n


# ---------------------------------------------------------------------------
# Dati del pezzo scritti DENTRO il DXF, letti dall'importatore DXF di Lantek
# ---------------------------------------------------------------------------
# Provato con Stefano il 06/10/2026: in Lantek Expert, importatore DXF,
# Configura -> Altri -> Testi: Quantita' "QTA", Materiale "MAT", Spessore "SP",
# Ordine "ORD" e spuntato "Importare le proprieta' dei testi". Una scritta
# "QTA 8" (etichetta, spazio, valore) dentro il contorno del pezzo compila la
# griglia di importazione. "QTA=8", "QTA:8", "QTA8" o etichetta e valore in
# due scritte separate NON funzionano.
LAYER_DATI = '0'


def _num_lantek(v, decimale: str = ',') -> str:
    s = f'{float(v):g}'
    return s.replace('.', decimale)


def _una_parola(v) -> str:
    """Lantek legge UNA parola dopo l'etichetta: gli spazi diventano "_"."""
    return '_'.join(str(v).split())


def testi_dati(riga: dict, commessa: str = '', decimale: str = ',',
               cliente: str = '', consegna=None) -> list:
    """Le scritte per un pezzo, nell'ordine dei campi di Altri -> Testi:
    ["QTA 4", "MAT FERRO", "SP 1,5", "ORD 1252", "CLI DECA_S.r.l.", "CONS 09/10/2026"].
    CLI e CONS vanno in User data 1 e 2 (Lantek non ha un testo per Cliente
    e Data consegna); ordine e cliente veri Stefano li mette dopo, in Lantek.
    Spessore: 1,5 e 1.5 li legge tutti e due (prova del 06/10/2026)."""
    out = [f"QTA {int(riga['quantita'])}"]
    if riga.get('materiale'):
        out.append(f"MAT {_una_parola(riga['materiale'])}")
    if riga.get('spessore') not in (None, ''):
        out.append(f"SP {_num_lantek(riga['spessore'], decimale)}")
    if commessa:
        out.append(f"ORD {_una_parola(commessa)}")
    if cliente:
        out.append(f"CLI {_una_parola(cliente)}")
    if consegna:
        d = consegna
        if isinstance(d, str):
            try:
                d = datetime.strptime(d[:10], '%Y-%m-%d')
            except ValueError:
                d = None
        if d:
            out.append(f"CONS {d.strftime('%d/%m/%Y')}")
    return out


def dxf_con_dati(sorgente: str, testi: list) -> bytes:
    """Il DXF `sorgente` con le scritte dei dati aggiunte dentro il contorno
    del pezzo (punto interno al contorno esterno; se il contorno non si trova,
    il centro del disegno). Il file d'origine non si tocca."""
    import io
    import ezdxf
    from ezdxf import bbox
    doc = ezdxf.readfile(sorgente)
    msp = doc.modelspace()
    e = bbox.extents(msp)
    cx = (e.extmin.x + e.extmax.x) / 2 if e.has_data else 0.0
    cy = (e.extmin.y + e.extmax.y) / 2 if e.has_data else 0.0
    lato = min(e.size.x, e.size.y) if e.has_data else 100.0
    try:
        from .preventivi.dxf_polygon_detector_v3 import contorno_pezzo_mm, scala_unita_mm
        outer = contorno_pezzo_mm(doc)
        if outer is not None:
            sc = scala_unita_mm(doc)[0] or 1.0
            p = outer.representative_point()
            cx, cy = p.x / sc, p.y / sc
            b = outer.bounds
            lato = min(b[2] - b[0], b[3] - b[1]) / sc
    except Exception:
        logger.debug('contorno per le scritte Lantek non trovato', exc_info=True)
    h = max(1.0, min(5.0, lato / 12))
    for i, t in enumerate(testi):
        msp.add_text(t, dxfattribs={'layer': LAYER_DATI, 'height': h,
                                    'insert': (cx, cy - i * h * 1.6)})
    buf = io.StringIO()
    doc.write(buf)
    return buf.getvalue().encode(doc.output_encoding or 'cp1252', errors='replace')


# ---------------------------------------------------------------------------
# Ordini di produzione per l'XML Importer di Lantek
# ---------------------------------------------------------------------------
# Provato con Stefano il 07/10/2026: C:\Lantek\System\Common\XmlImporter.exe
# (programma di Lantek, coperto dalla licenza) con un comando MANUFACTURING
# crea l'ordine di produzione con quantita', ordine, cliente e consegna
# (FT-PROVA-03: 7 pezzi, uguale a uno fatto a mano). E' Lantek a scrivere,
# quando Stefano lancia l'import: FerroTrack prepara solo il file.
# I pezzi NUOVI invece l'XML li crea come materia prima (PType 1): i disegni
# si importano ancora dal MES ("Importa -> Files DXF"), poi questo file.
CENTRO_LAVORO = 'CY Laser 3015 HL ECS'
OPERAZIONE = '2D Cut'
_FORME_SOCIETA = {'SRL', 'SPA', 'SNC', 'SAS', 'SRLS', 'SS', 'SC', 'SOCIETA'}


def ordini_in_lantek(commessa: str) -> dict | None:
    """Gli ordini di produzione che Lantek ha gia' per quel numero d'ordine
    (OrdRef "1252", anche "1252 C75"): {codice in maiuscolo: quantita' totale}.
    None se Lantek non e' leggibile."""
    commessa = str(commessa or '').strip()
    cfg = _config()
    if not commessa or not cfg.get('abilitato'):
        return {} if commessa == '' else None
    chiave = ('ordini', cfg['server'], cfg['database'], commessa)
    with _LOCK:
        c = _CACHE.get(chiave)
        if c and time.time() - c[0] < _DURATA_CACHE_S:
            return c[1]
    like = commessa.replace('[', '[[]').replace('%', '[%]').replace('_', '[_]') + ' %'
    try:
        con, cur = _connessione(cfg)
        try:
            cur.execute('SELECT UPPER(LTRIM(RTRIM(PrdRef))), SUM(Quantity) FROM MMNN_MMOO_00000100 '
                        'WHERE LTRIM(RTRIM(OrdRef)) = ? OR LTRIM(OrdRef) LIKE ? '
                        'GROUP BY UPPER(LTRIM(RTRIM(PrdRef)))', [commessa, like])
            out = {str(p): float(q or 0) for p, q in cur.fetchall()}
        finally:
            con.close()
    except Exception as e:
        logger.info('ordini di produzione Lantek non letti: %s', e)
        return None
    with _LOCK:
        _CACHE[chiave] = (time.time(), out)
    return out


def segna_gia_ordinati(righe: list, ordinati: dict | None) -> None:
    """Su ogni riga: 'in_produzione' = pezzi che Lantek ha gia' negli ordini
    di produzione di quell'ordine (0 se nessuno), con l'avviso se la quantita'
    e' diversa da quella dell'ordine."""
    for r in righe:
        q = (ordinati or {}).get(str(r['codice']).strip().upper(), 0)
        r['in_produzione'] = int(q) if float(q).is_integer() else q
        if q and abs(q - r['quantita']) > 0.001:
            r['avvisi'].append(f"in Lantek in produzione {r['in_produzione']:g} pz, l'ordine ne chiede {r['quantita']}")


def clienti_lantek() -> list:
    """Le sigle dei clienti usate negli ordini di produzione di Lantek
    (CusRef, ultimi due anni), le piu' usate prima. [] senza Lantek."""
    cfg = _config()
    if not cfg.get('abilitato'):
        return []
    chiave = ('clienti', cfg['server'], cfg['database'])
    with _LOCK:
        c = _CACHE.get(chiave)
        if c and time.time() - c[0] < 600:
            return c[1]
    try:
        con, cur = _connessione(cfg)
        try:
            cur.execute("SELECT UPPER(LTRIM(RTRIM(CusRef))) c, COUNT(*) n FROM MMNN_MMOO_00000100 "
                        "WHERE CrtDate >= DATEADD(year, -2, GETDATE()) AND LTRIM(RTRIM(CusRef)) <> '' "
                        "GROUP BY UPPER(LTRIM(RTRIM(CusRef))) ORDER BY n DESC")
            out = [str(r[0]) for r in cur.fetchall()]
        finally:
            con.close()
    except Exception as e:
        logger.info('clienti Lantek non letti: %s', e)
        return []
    with _LOCK:
        _CACHE[chiave] = (time.time(), out)
    return out


def cliente_lantek(nome: str, noti: list | None = None) -> str:
    """La sigla del cliente come la si scrive in Lantek: "DECA S.r.l." -> "DECA",
    "B&B Italia S.p.A." -> "B&B". Se Lantek usa gia' una sigla che corrisponde
    all'inizio del nome si usa quella (niente "POLIFOMR"), altrimenti la prima
    parola del nome."""
    # "S.r.l." -> "SRL": i punti si tolgono senza spezzare la parola
    s = re.sub(r'[^\w&\s]', ' ', str(nome or '').upper().replace('.', ''))
    parole = [p for p in s.split() if p not in _FORME_SOCIETA]
    if not parole:
        return ''
    pulito = ' '.join(parole)
    candidati = [n for n in (noti or []) if n and (pulito == n or pulito.startswith(n + ' '))]
    if candidati:
        return max(candidati, key=len)
    return parole[0][:40]


def riferimento_ordine(commessa: str, codice: str) -> str:
    """Il riferimento dell'ordine di produzione (max 40 caratteri in Lantek):
    sempre lo stesso per ordine+codice, cosi' reimportando il file Lantek
    aggiorna l'ordine invece di farne un doppione."""
    rif = f'FT{commessa}-{codice}' if commessa else f'FT-{codice}'
    if len(rif) <= 40:
        return rif
    import hashlib
    return rif[:31] + '-' + hashlib.sha1(rif.encode('utf-8')).hexdigest()[:8]


def righe_per_xml(righe: list) -> list:
    """Le righe che Lantek puo' importare: il pezzo deve esserci gia' (stato
    "in_lantek"), o Lantek non era leggibile ("sconosciuto": ci si prova).
    Fuori quelle che hanno gia' un ordine di produzione per quell'ordine: si
    farebbe un doppione."""
    return [r for r in righe if r.get('stato') != 'nuovo' and int(r.get('quantita') or 0) > 0
            and not r.get('in_produzione')]


def xml_ordini_produzione(righe: list, commessa: str = '', cliente: str = '', consegna=None,
                          centro: str = CENTRO_LAVORO, operazione: str = OPERAZIONE) -> bytes:
    """Il file per l'XML Importer: un comando MANUFACTURING per riga, nel
    formato dell'esempio di Lantek (C:\\Lantek\\System\\Masterlink\\Samples\\
    Import.xml). FldType: 20 testo, 100 numero, 120 data AAAAMMGG."""
    from xml.sax.saxutils import quoteattr
    if isinstance(consegna, str):
        try:
            consegna = datetime.strptime(consegna[:10], '%Y-%m-%d')
        except ValueError:
            consegna = None

    def campo(ref, val, tipo):
        return f'\t\t<FIELD FldRef="{ref}" FldValue={quoteattr(str(val))} FldType="{tipo}" />\n'

    out = ['<?xml version="1.0" encoding="utf-8"?>\n<DATAEX>\n']
    for r in righe:
        out.append('\t<COMMAND Name="Import" TblRef="MANUFACTURING">\n')
        out.append(campo('Reference', riferimento_ordine(commessa, r['codice']), 20))
        out.append(campo('Product', r['codice'], 20))
        out.append(campo('WorkCenter', centro, 20))
        out.append(campo('Operation', operazione, 20))
        if cliente:
            out.append(campo('Customer', cliente[:40], 20))
        if commessa:
            out.append(campo('SaleOrder', commessa[:40], 20))
        out.append(campo('Quantity', int(r['quantita']), 100))
        if consegna:
            out.append(campo('DeliveryDate', consegna.strftime('%Y%m%d'), 120))
        out.append('\t</COMMAND>\n')
    out.append('</DATAEX>\n')
    return ''.join(out).encode('utf-8')


# ---------------------------------------------------------------------------
# Lancio dell'XML Importer ("Manda a Lantek", dopo la conferma di Stefano)
# ---------------------------------------------------------------------------
# Argomenti a coppie "nome valore" (letti dal programma il 07/10/2026 e provati
# con un file sonda): -src <file> -HIDE 1 -CloseWindow 1 -Showlog 0. Senza
# finestra; il rapporto va in "<cartella del file>\Importer log\
# <nome>_<data ora>_log.html" (o _logERR.html se ci sono errori), e l'XML
# viene spostato li'. A scrivere in Lantek e' il programma di Lantek.
XMLIMPORTER = r'C:\Lantek\System\Common\XmlImporter.exe'
_LOCK_IMPORT = threading.Lock()


def _xmlimporter() -> str:
    return (_config().get('xmlimporter') or XMLIMPORTER)


def xmlimporter_disponibile() -> bool:
    import os
    return os.name == 'nt' and os.path.isfile(_xmlimporter())


def leggi_rapporto(percorso: str) -> dict:
    """Il rapporto HTML dell'XML Importer: totali e, per ogni comando con
    errore, il pezzo e il messaggio."""
    import html as _html
    with open(percorso, encoding='utf-8', errors='replace') as f:
        testo = f.read()
    righe = [r.strip() for r in _html.unescape(re.sub(r'<[^>]+>', '\n', testo)).splitlines()]
    righe = [r for r in righe if r and not r.endswith(';') and r not in ('">', '"')]

    def numero(etichetta):
        # "Total commands:" e il numero possono stare nella stessa cella o in due
        for i, r in enumerate(righe):
            if r.lower().startswith(etichetta.lower()):
                resto = r[len(etichetta):].strip(' :\t')
                if not resto and i + 1 < len(righe):
                    resto = righe[i + 1]
                m = re.match(r'\d+', resto)
                return int(m.group(0)) if m else None
        return None

    # un blocco per comando: "1. IMPORT FOR TABLE ..." fino al successivo
    blocchi, attuale = [], None
    for r in righe:
        if re.match(r'\d+\.\s+IMPORT FOR TABLE', r):
            attuale = [r]
            blocchi.append(attuale)
        elif attuale is not None:
            attuale.append(r)
    errori = []
    for b in blocchi:
        # riuscito: "successful", oppure nessun messaggio (IMPORTGEO riuscito)
        if len(b) <= 2 or any(re.search(r'\bsuccessful\b', r, re.I) for r in b[1:]):
            continue
        testa = ' '.join(b[1:2])
        mp = re.search(r'Destination Product:\s*(.+)$', testa) or re.search(r'Product:\s*(\S+)', testa)
        msg = b[-1] if len(b) > 2 else ''
        if 'BROKENRULE' in msg:
            msg = re.sub(r'.*Text="([^"]*)".*', r'\1', msg)
        errori.append({'comando': b[0], 'pezzo': mp.group(1).strip() if mp else None, 'messaggio': msg})
    return {'totale': numero('Total commands'), 'ok': numero('Commands Ok'),
            'avvisi': numero('Commands with warning'), 'con_errore': numero('Commands with error'),
            'errori': errori}


def importa_xml(contenuto: bytes, nome: str, cartella: str, attesa_s: int = 300) -> dict:
    """Salva l'XML in `cartella` e lo fa importare all'XML Importer di Lantek
    (senza finestra), poi legge il suo rapporto. Uno alla volta.

    {'eseguito': bool, 'errore': str|None, 'rapporto': {...}, 'file_rapporto': str}"""
    import glob
    import os
    import subprocess
    exe = _xmlimporter()
    if not xmlimporter_disponibile():
        return {'eseguito': False, 'errore': f'XML Importer di Lantek non trovato ({exe})'}
    if not _LOCK_IMPORT.acquire(blocking=False):
        return {'eseguito': False, 'errore': 'Un altro invio a Lantek è in corso: riprova tra poco'}
    try:
        # aperto a mano da qualcuno? Con -HIDE il secondo si chiuderebbe senza importare
        try:
            gia = subprocess.run(['tasklist', '/FI', 'IMAGENAME eq XmlImporter.exe', '/NH'],
                                 capture_output=True, text=True, timeout=15,
                                 creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
            if 'xmlimporter.exe' in (gia.stdout or '').lower():
                return {'eseguito': False,
                        'errore': "XML Importer è già aperto su questo PC: chiudilo e riprova"}
        except Exception:
            logger.debug('controllo XmlImporter aperto non riuscito', exc_info=True)
        os.makedirs(cartella, exist_ok=True)
        base = re.sub(r'[^\w\-]+', '_', nome).strip('_')[:60] or 'ordini'
        base = f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        percorso = os.path.join(cartella, base + '.xml')
        with open(percorso, 'wb') as f:
            f.write(contenuto)
        try:
            p = subprocess.run([exe, '-src', percorso, '-HIDE', '1', '-CloseWindow', '1', '-Showlog', '0'],
                               cwd=os.path.dirname(exe), timeout=attesa_s, capture_output=True,
                               creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        except subprocess.TimeoutExpired:
            return {'eseguito': False, 'errore': f'XML Importer non ha finito in {attesa_s} s: '
                                                 'controlla in Lantek prima di riprovare'}
        rapporti = sorted(glob.glob(os.path.join(cartella, 'Importer log', glob.escape(base) + '_*log*.html')),
                          key=os.path.getmtime)
        if not rapporti:
            return {'eseguito': False, 'errore': 'XML Importer non ha lasciato il rapporto '
                                                 f'(uscita {p.returncode}): controlla in Lantek'}
        return {'eseguito': True, 'errore': None, 'rapporto': leggi_rapporto(rapporti[-1]),
                'file_rapporto': rapporti[-1]}
    finally:
        _LOCK_IMPORT.release()
        # gli ordini di produzione sono cambiati: si rilegge Lantek
        with _LOCK:
            for k in [k for k in _CACHE if k and k[0] == 'ordini']:
                _CACHE.pop(k, None)


def scrivi_excel(percorso_o_file, righe: list, data_consegna=None, commessa: str = '') -> None:
    """Il file nel formato del modello iErp "ImportProduzione" (Foglio1, prima
    riga d'intestazione, una riga per codice)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = 'Foglio1'
    ws.append(INTESTAZIONE)
    for cella in ws[1]:
        cella.font = Font(bold=True)
    if isinstance(data_consegna, str):
        try:
            data_consegna = datetime.strptime(data_consegna[:10], '%Y-%m-%d')
        except ValueError:
            data_consegna = None
    elif isinstance(data_consegna, date) and not isinstance(data_consegna, datetime):
        data_consegna = datetime.combine(data_consegna, datetime.min.time())
    for r in righe:
        ws.append([r['codice'], int(r['quantita']), r['materiale'] or '',
                   r['spessore'] if r['spessore'] is not None else None,
                   data_consegna, commessa or ''])
    for riga in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for cella in riga:
            cella.number_format = 'DD/MM/YYYY'
    for col, larg in zip('ABCDEF', (28, 10, 14, 10, 15, 16)):
        ws.column_dimensions[col].width = larg
    wb.save(percorso_o_file)
