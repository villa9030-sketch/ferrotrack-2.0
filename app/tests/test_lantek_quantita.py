"""Quantita' per Lantek: Excel (iErp), scritte nei DXF, XML degli ordini di produzione.

Copre:
 1. dalla distinta: solo lamiere col codice, quantita' sommate (lotto, assiemi)
 2. con Lantek: materiale e spessore come in Lantek per i codici che ci sono;
    avvisi per codici nuovi, materiale/spessore diversi, revisione piu' nuova
 3. senza Lantek (spento): il file si fa lo stesso coi dati di FerroTrack
 4. il file: Foglio1, intestazione del modello ImportProduzione di iErp,
    data di consegna come data, commessa = numero dell'ordine
 5. API: elenco coi controlli e download per laser e ufficio, non per l'officina
 6. nella cartella dell'ordine niente Excel; 7. scritte nei DXF; 8. XML (XmlImporter)

Lantek e' SIMULATO: i test non toccano il database vero di Lantek.
Gira su DATABASE TEMPORANEO. Esecuzione: python app/tests/test_lantek_quantita.py
"""
import io
import json
import os
import shutil
import sys
import tempfile
import uuid
from datetime import datetime, timedelta

_QUI = os.path.dirname(os.path.abspath(__file__))
_APP = os.path.dirname(_QUI)
if _APP not in sys.path:
    sys.path.insert(0, _APP)
os.environ['FERROTRACK_SKIP_DB_INIT'] = '1'

import werkzeug  # noqa: E402
if not hasattr(werkzeug, '__version__'):
    werkzeug.__version__ = '3.0.0'

from sqlalchemy import create_engine  # noqa: E402

from backend import models  # noqa: E402
from backend.models import Base, User, Order, Preventivo  # noqa: E402
from backend import models_ore  # noqa: E402,F401

_TMP = os.path.join(tempfile.gettempdir(), f'test_ltq_{uuid.uuid4().hex[:8]}.db')
_ENG = create_engine('sqlite:///' + _TMP.replace(chr(92), '/'), connect_args={'check_same_thread': False})
models.engine = _ENG
models.SessionLocal.configure(bind=_ENG)
Base.metadata.create_all(bind=_ENG)

import importlib  # noqa: E402
A = importlib.import_module('backend.app')
from backend import lantek as L  # noqa: E402
from backend.database import PreventivoManager  # noqa: E402

_CARTELLE = tempfile.mkdtemp(prefix='ft_ltq_')
A.UPLOAD_FOLDER = os.path.join(_CARTELLE, 'uploads')
A.DRAWINGS_FOLDER = os.path.join(A.UPLOAD_FOLDER, 'drawings')
os.makedirs(A.DRAWINGS_FOLDER, exist_ok=True)
_RETE = os.path.join(_CARTELLE, 'rete')
_cfg_vera = A.ConfigManager.load_config
A.ConfigManager.load_config = staticmethod(lambda: {**(_cfg_vera() or {}), 'disegni_export_root': _RETE})

OK = 0
KO = []


def check(nome, cond, extra=''):
    global OK
    if cond:
        OK += 1
        print(f'  [OK] {nome}')
    else:
        KO.append(nome)
        print(f'  [KO] {nome}  {extra}')


LANTEK_FINTO = {'disponibile': True, 'errore': None, 'pezzi': {
    'A1-00': {'esiste': True, 'materiale': 'FERRO', 'spessore': 3.0, 'revisioni': []},
    'B1-00': {'esiste': True, 'materiale': 'INOX', 'spessore': 1.5, 'revisioni': ['B1-01']},
    'C1-00': {'esiste': False, 'materiale': None, 'spessore': None, 'revisioni': []},
}}


def preventivo_accettato(cliente, numero_cliente):
    pr = PreventivoManager.create(cliente, 'paolo', quantita=2, numero_ordine_cliente=numero_cliente)
    pid = pr['id'] if isinstance(pr, dict) else pr
    PreventivoManager.replace_articoli(pid, [
        {'codice': 'A1-00', 'quantita': 3, 'materiale': 'S235', 'spessore_mm': 3, 'costo_base_override': 5},
        {'codice': 'B1-00', 'quantita': 2, 'materiale': 'INOX_304', 'spessore_mm': 2,
         'codice_assieme': 'ASS1', 'costo_base_override': 4},
        {'codice': 'C1-00', 'quantita': 1, 'materiale': 'ALU', 'spessore_mm': 4, 'costo_base_override': 3},
    ])
    PreventivoManager.replace_assiemi(pid, [{'codice_assieme': 'ASS1', 'qty': 4, 'costo': 10}])
    PreventivoManager.replace_tubolari(pid, [{'profilo': 'Quadro 40x40 sp.2mm', 'lunghezza_m': 0.5, 'qty': 2}])
    s = models.SessionLocal()
    try:
        s.query(Preventivo).filter(Preventivo.id == pid).first().status = 'INVIATO'
        s.commit()
    finally:
        s.close()
    return PreventivoManager.accetta_e_crea_ordine(pid, 'paolo').get('order_id')


def main():
    s = models.SessionLocal()
    s.add(User(id='paolo', name='Paolo Capo', role='Capo Officina', is_active=True, is_capo=True))
    s.commit(); s.close()

    print('1) Righe dalla distinta')
    distinta = [
        {'tipo': 'lamiera', 'codice': 'A1-00', 'quantita': 6, 'materiale': 'S235', 'spessore_mm': 3},
        {'tipo': 'lamiera', 'codice': 'A1-00', 'quantita': 2, 'materiale': 'S235', 'spessore_mm': 3},
        {'tipo': 'lamiera', 'codice': 'B1-00', 'quantita': 8, 'materiale': 'INOX_304', 'spessore_mm': 2},
        {'tipo': 'lamiera', 'codice': 'C1-00', 'quantita': 2, 'materiale': 'ALU', 'spessore_mm': 4},
        {'tipo': 'tubolare', 'codice': 'Quadro', 'quantita': 2},
        {'tipo': 'assieme', 'codice': 'ASS1', 'quantita': 4},
        {'tipo': 'lamiera', 'codice': '', 'quantita': 1},
    ]
    r = {x['codice']: x for x in L.righe_quantita(distinta, LANTEK_FINTO)}
    check('solo lamiere col codice', sorted(r) == ['A1-00', 'B1-00', 'C1-00'], sorted(r))
    check('quantita sommate per codice (6 + 2 = 8)', r['A1-00']['quantita'] == 8, r['A1-00'])

    print('\n2) Con Lantek')
    check('A1: in Lantek, FERRO 3, nessun avviso', r['A1-00']['stato'] == 'in_lantek'
          and r['A1-00']['materiale'] == 'FERRO' and r['A1-00']['spessore'] == 3.0 and not r['A1-00']['avvisi'], r['A1-00'])
    b = r['B1-00']
    check('B1: si usa lo spessore di Lantek (1,5) e si avvisa (ordine 2)',
          b['spessore'] == 1.5 and any('spessore' in a for a in b['avvisi']), b)
    check('B1: avviso revisione piu nuova (B1-01)', any('B1-01' in a for a in b['avvisi']), b['avvisi'])
    c = r['C1-00']
    check('C1: nuovo, materiale ALLUMINIO dalla tabella', c['stato'] == 'nuovo' and c['materiale'] == 'ALLUMINIO'
          and any('importa prima il disegno' in a for a in c['avvisi']), c)

    print('\n2b) Codice scritto senza revisione, commessa')
    finto = {'disponibile': True, 'errore': None, 'pezzi': {
        'D1': {'esiste': True, 'codice_lantek': 'D1-00', 'materiale': 'FERRO', 'spessore': 2.0, 'revisioni': ['D1-00']},
        'E1': {'esiste': False, 'codice_lantek': None, 'materiale': None, 'spessore': None, 'revisioni': ['E1-00', 'E1-01']},
    }}
    rr = {x['codice_ft']: x for x in L.righe_quantita([
        {'tipo': 'lamiera', 'codice': 'D1', 'quantita': 2, 'materiale': 'S235', 'spessore_mm': 2},
        {'tipo': 'lamiera', 'codice': 'E1', 'quantita': 1, 'materiale': 'S235', 'spessore_mm': 2}], finto)}
    check('D1 -> nel file il codice di Lantek D1-00', rr['D1']['codice'] == 'D1-00' and rr['D1']['stato'] == 'in_lantek'
          and not rr['D1']['avvisi'], rr['D1'])
    check('E1 con due revisioni in Lantek: non si sceglie, si avvisa',
          rr['E1']['codice'] == 'E1' and any('scegli' in a for a in rr['E1']['avvisi']), rr['E1'])
    check('commessa: "A 001252" -> "1252", "061" resta, PREV resta',
          (L.commessa_lantek('A 001252'), L.commessa_lantek('061'), L.commessa_lantek('PREV-2026-0005'))
          == ('1252', '061', 'PREV-2026-0005'))

    print('\n3) Senza Lantek')
    r2 = {x['codice']: x for x in L.righe_quantita(distinta, {'disponibile': False, 'pezzi': {}})}
    check('il file si fa coi dati di FerroTrack', r2['B1-00']['materiale'] == 'INOX' and r2['B1-00']['spessore'] == 2.0
          and r2['B1-00']['stato'] == 'sconosciuto' and not r2['B1-00']['avvisi'], r2['B1-00'])

    print('\n4) Il file Excel')
    buf = io.BytesIO()
    L.scrivi_excel(buf, list(r.values()), '2030-11-06', '1252')
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(buf.getvalue())).active
    check('foglio "Foglio1"', ws.title == 'Foglio1')
    intest = [c.value for c in ws[1]]
    check('intestazione del modello iErp', intest == ['Codice Articolo', 'Quantità', 'Materiale', 'Spessore',
                                                      'Data consegna', 'Commessa'], intest)
    riga = [c.value for c in ws[2]]
    check('riga: codice, quantita, materiale, spessore, data, commessa',
          riga[0] == 'A1-00' and riga[1] == 8 and riga[2] == 'FERRO' and riga[3] == 3.0
          and isinstance(riga[4], datetime) and riga[4].date().isoformat() == '2030-11-06' and riga[5] == '1252', riga)

    print('\n5) API')
    from tests.accesso_aiuto import postazioni, persona, entra, modalita
    modalita('protetto')
    postazioni()
    persona('ufficio', 'Elena Ufficio', 'Amministrazione', pin='2580')
    uff = entra(A.app, 'ufficio', pin='2580')
    laser = entra(A.app, 'laser')
    rep = entra(A.app, 'reparto')
    L.pezzi_in_lantek = lambda codici: {**LANTEK_FINTO, 'pezzi': {k: v for k, v in LANTEK_FINTO['pezzi'].items() if k in codici}}
    L.clienti_lantek = lambda: ['DECA', 'B&B', 'POLIFORM']
    # B1-00 ha gia' un ordine di produzione in Lantek per il 1252 (8 pezzi)
    L.ordini_in_lantek = lambda commessa: {'B1-00': 8.0} if commessa == '1252' else {}
    L.ordini_fatti_in_lantek = lambda commessa: {}
    oid = preventivo_accettato('DECA S.r.l.', '1252')
    check('ordine creato', bool(oid))
    s = models.SessionLocal()
    o = s.query(Order).filter(Order.id == oid).first()
    o.data_consegna = datetime(2030, 11, 6)
    s.commit(); s.close()
    d = laser.get(f'/api/orders/{oid}/lantek-quantita').get_json() or {}
    check('il laser vede le righe coi controlli', d.get('success') and d.get('n_codici') == 3
          and d.get('n_nuovi') == 1 and d.get('commessa') == '1252', {k: d.get(k) for k in ('n_codici', 'n_nuovi', 'commessa', 'error')})
    q = {x['codice']: x['quantita'] for x in d.get('righe') or []}
    check('quantita: A1 3x2=6, B1 2x4=8, C1 1x2=2', q == {'A1-00': 6, 'B1-00': 8, 'C1-00': 2}, q)
    rx = uff.get(f'/api/orders/{oid}/lantek-quantita.xlsx')
    check('l\'ufficio scarica l\'Excel', rx.status_code == 200 and rx.data[:2] == b'PK', rx.status_code)
    check('nome del file col numero del cliente', '1252' in (rx.headers.get('Content-Disposition') or ''),
          rx.headers.get('Content-Disposition'))
    check('il tablet officina no', rep.get(f'/api/orders/{oid}/lantek-quantita').status_code in (401, 403))
    check('cliente come in Lantek: DECA', d.get('cliente_lantek') == 'DECA', d.get('cliente_lantek'))
    check('B1 gia\' in produzione in Lantek: contato, nel file ne resta 1',
          d.get('n_in_produzione') == 1 and d.get('n_xml') == 1, {k: d.get(k) for k in ('n_in_produzione', 'n_xml')})
    rx = laser.get(f'/api/orders/{oid}/lantek-ordini.xml')
    check('il laser scarica l\'XML degli ordini di produzione', rx.status_code == 200
          and 'ORDINI LANTEK - 1252.xml' in (rx.headers.get('Content-Disposition') or ''),
          (rx.status_code, rx.headers.get('Content-Disposition')))
    import xml.etree.ElementTree as ET
    cmds = ET.fromstring(rx.data).findall('COMMAND')
    per = {c.find("FIELD[@FldRef='Product']").get('FldValue'): {f.get('FldRef'): f.get('FldValue') for f in c}
           for c in cmds}
    check('solo i codici gia\' in Lantek e non ancora in produzione (C1 nuovo e B1 esclusi)',
          sorted(per) == ['A1-00'], sorted(per))
    a = per.get('A1-00') or {}
    check('A1: quantita 6, ordine 1252, cliente DECA, consegna 20301106, CY Laser, 2D Cut',
          a.get('Quantity') == '6' and a.get('SaleOrder') == '1252' and a.get('Customer') == 'DECA'
          and a.get('DeliveryDate') == '20301106' and a.get('WorkCenter') == 'CY Laser 3015 HL ECS'
          and a.get('Operation') == '2D Cut' and a.get('Reference') == 'FT1252-A1-00', a)
    check('l\'officina no', rep.get(f'/api/orders/{oid}/lantek-ordini.xml').status_code in (401, 403))

    print('\n6) Nella cartella dell\'ordine')
    os.makedirs(os.path.join(A.DRAWINGS_FOLDER, oid), exist_ok=True)
    import ezdxf as _ez
    _d = _ez.new()
    _d.modelspace().add_lwpolyline([(0, 0), (100, 0), (100, 50), (0, 50)], close=True)
    _d.saveas(os.path.join(A.DRAWINGS_FOLDER, oid, 'A1-00.dxf'))
    e = A._esporta_disegni_per_officina(oid)
    fatti = [f for _b, _d2, fs in os.walk(e.get('percorso') or '') for f in fs]
    check('nella cartella niente Excel delle quantita',
          not any(f.endswith('.xlsx') for f in fatti) and 'quantita_lantek' not in e, (e, fatti))

    print('\n7) Scritte dentro i DXF per l\'importatore di Lantek')
    t = L.testi_dati({'quantita': 4, 'materiale': 'FERRO', 'spessore': 1.5}, '1252', ',', 'DECA S.r.l.', '2030-11-06')
    check('formato provato in Lantek: etichetta, spazio, una parola',
          t == ['QTA 4', 'MAT FERRO', 'SP 1,5', 'ORD 1252', 'CLI DECA_S.r.l.', 'CONS 06/11/2030'], t)
    import ezdxf
    from shapely.geometry import Point, Polygon
    pdxf = os.path.join(_CARTELLE, 'pezzo.dxf')
    doc = ezdxf.new(); doc.header['$INSUNITS'] = 4
    # pezzo a L: il centro dell'ingombro cade FUORI dal pezzo
    forma = [(0, 0), (300, 0), (300, 40), (40, 40), (40, 300), (0, 300)]
    doc.modelspace().add_lwpolyline(forma, close=True)
    doc.saveas(pdxf)
    prima = open(pdxf, 'rb').read()
    out = L.dxf_con_dati(pdxf, t)
    import io as _io
    d2 = ezdxf.read(_io.StringIO(out.decode('utf-8', 'replace')))
    testi = [e for e in d2.modelspace().query('TEXT')]
    check('le scritte ci sono', [e.dxf.text for e in testi] == t, [e.dxf.text for e in testi])
    check('la prima scritta e\' dentro il pezzo (anche a L)',
          Polygon(forma).contains(Point(testi[0].dxf.insert[0], testi[0].dxf.insert[1])), testi[0].dxf.insert)
    check('il file d\'origine non si tocca', open(pdxf, 'rb').read() == prima)
    dati = {'25ab1979': ('25AB1979-00', ['QTA 2'])}
    check('nome del file col codice di Lantek', A._voce_per_lantek('X/S235 - 3 mm/25AB1979.dxf', dati)
          == ('X/S235 - 3 mm/25AB1979-00.dxf', ['QTA 2']))
    check('originali e da preparare non si toccano',
          A._voce_per_lantek('X/_DISEGNI ORIGINALI/25AB1979.dxf', dati) is None
          and A._voce_per_lantek('X/_DA PREPARARE/S235 - 3 mm/25AB1979.dxf', dati) is None)

    print('\n8) XML per l\'XML Importer di Lantek')
    noti = ['DECA', 'B&B', 'POLIFORM', 'AZA']
    check('cliente: sigla gia\' usata in Lantek', (L.cliente_lantek('DECA S.r.l.', noti),
          L.cliente_lantek('B&B Italia S.p.A.', noti), L.cliente_lantek('Poliform', noti))
          == ('DECA', 'B&B', 'POLIFORM'))
    check('cliente nuovo: prima parola senza forma societaria',
          L.cliente_lantek('S.r.l. Rossi Lamiere', []) == 'ROSSI' and L.cliente_lantek('', noti) == '')
    rif = L.riferimento_ordine('1252', 'X' * 60)
    check('riferimento sempre uguale e al massimo 40 caratteri',
          len(rif) == 40 and rif == L.riferimento_ordine('1252', 'X' * 60)
          and L.riferimento_ordine('1252', 'A1-00') == 'FT1252-A1-00', rif)
    righe = [{'codice': 'B&B DADO 244', 'quantita': 40, 'stato': 'in_lantek'},
             {'codice': 'NUOVO-00', 'quantita': 2, 'stato': 'nuovo'},
             {'codice': 'SENZA-00', 'quantita': 3, 'stato': 'sconosciuto'}]
    sel = L.righe_per_xml(righe)
    check('i nuovi esclusi, senza Lantek ci si prova', [r['codice'] for r in sel] == ['B&B DADO 244', 'SENZA-00'])
    x = L.xml_ordini_produzione(sel, '1252', 'B&B', None)
    radice = ET.fromstring(x)
    c0 = {f.get('FldRef'): (f.get('FldValue'), f.get('FldType')) for f in radice.find('COMMAND')}
    check('formato di Lantek: DATAEX / COMMAND Import MANUFACTURING / FIELD',
          radice.tag == 'DATAEX' and all(c.get('Name') == 'Import' and c.get('TblRef') == 'MANUFACTURING'
                                         for c in radice.findall('COMMAND')))
    check('"&" nel codice e nel cliente scritto bene, quantita numero (100), senza consegna niente data',
          c0['Product'] == ('B&B DADO 244', '20') and c0['Customer'] == ('B&B', '20')
          and c0['Quantity'] == ('40', '100') and 'DeliveryDate' not in c0, c0)

    rr = [{'codice': 'K1-00', 'quantita': 6, 'avvisi': []}, {'codice': 'K2-00', 'quantita': 4, 'avvisi': []},
          {'codice': 'K3-00', 'quantita': 2, 'avvisi': []}]
    L.segna_gia_ordinati(rr, {'K1-00': 6.0, 'K2-00': 1.0})
    check('gia\' in produzione: uguale nessun avviso, diversa (1 invece di 4) avviso, assente 0',
          rr[0]['in_produzione'] == 6 and not rr[0]['avvisi'] and rr[1]['in_produzione'] == 1
          and any('ne chiede 4' in a for a in rr[1]['avvisi']) and rr[2]['in_produzione'] == 0, rr)
    check('quelli in produzione non vanno nel file',
          [r['codice'] for r in L.righe_per_xml([{**r, 'stato': 'in_lantek'} for r in rr])] == ['K3-00'])
    # ordine 1030: tagliato ad agosto (ordini FATTI) e da rifare -> si rimanda
    rr = [{'codice': 'K4-00', 'quantita': 3, 'avvisi': [], 'stato': 'in_lantek'}]
    L.segna_gia_ordinati(rr, {}, {'K4-00': (3.0, '2026-08-06')})
    check('gia\' fatto in passato: segnato ma va nel file (si rifà)',
          rr[0]['in_produzione'] == 0 and rr[0]['gia_fatti'] == 3 and rr[0]['gia_fatti_il'] == '2026-08-06'
          and [r['codice'] for r in L.righe_per_xml(rr)] == ['K4-00'], rr)
    check('riferimento con l\'invio: mai uguale a uno vecchio',
          L.riferimento_ordine('1030', 'K4-00', '2610071010') == 'FT1030-K4-00-2610071010'
          and b'FT1030-K4-00-2610071010' in L.xml_ordini_produzione(rr, '1030', 'DECA', None, invio='2610071010'))

    print('\n9) "Manda a Lantek": rapporto dell\'XML Importer e invio')
    rapporto = os.path.join(_CARTELLE, 'r_logERR.html')
    with open(rapporto, 'w', encoding='utf-8') as f:
        f.write('''<html><body><table>
 <tr><td>Total commands:</td> <td>2</td></tr> <tr><td>Commands Ok:</td> <td>1</td></tr>
 <tr><td>Commands with warning:</td> <td>0</td></tr> <tr><td>Commands with error:</td> <td>1</td></tr></table>
 <div><div>
1. IMPORT FOR TABLE MANUFACTURING <div>
 Production Order: FT1252-A1-00 Destination Product: A1-00  </div></div>
 <ul style= " list-style-type: circle;
margin-top: 5px;
"> <li style="  color: #2F738C;
"><div> Insert into table 'MANUFACTURING' was successful.</div> </li></ul></div>
 <div><div>
2. IMPORT FOR TABLE MANUFACTURING <div>
 Production Order: FT1252-Z9-00 Destination Product: Z9-00  </div></div>
 <ul style= " list-style-type: circle;
"> <li style=" color: #E55B3C;
"> Cannot find product with reference: 'Z9-00' - insert failed.</li></ul></div>
</body></html>''')
    rp = L.leggi_rapporto(rapporto)
    check('rapporto: totali (2, 1 ok, 1 errore) e l\'errore col pezzo',
          (rp['totale'], rp['ok'], rp['avvisi'], rp['con_errore']) == (2, 1, 0, 1)
          and len(rp['errori']) == 1 and rp['errori'][0]['pezzo'] == 'Z9-00'
          and 'Cannot find product' in rp['errori'][0]['messaggio'], rp)

    # l'XML Importer e' SIMULATO: nessun programma di Lantek viene lanciato
    lanciati = []

    def importa_finto(contenuto, nome, cartella, attesa_s=300):
        lanciati.append((contenuto, nome, cartella))
        L.ordini_in_lantek = lambda commessa: {'B1-00': 8.0, 'A1-00': 6.0}
        return {'eseguito': True, 'errore': None, 'file_rapporto': 'x',
                'rapporto': {'totale': 1, 'ok': 1, 'avvisi': 0, 'con_errore': 0, 'errori': []}}
    L.importa_xml = importa_finto
    L.xmlimporter_disponibile = lambda: True
    L.procesos_disponibile = lambda: False
    L.ordini_in_lantek = lambda commessa: {'B1-00': 8.0} if commessa == '1252' else {}
    d = laser.get(f'/api/orders/{oid}/lantek-quantita').get_json() or {}
    check('per la conferma: elenco da mandare (A1 x 6) e invio automatico possibile',
          [(x['codice'], x['quantita']) for x in d.get('da_mandare') or []] == [('A1-00', 6)]
          and d.get('invio_automatico') is True,
          {k: d.get(k) for k in ('da_mandare', 'invio_automatico')})
    r9 = laser.post(f'/api/orders/{oid}/lantek-invia', json={'codici': ['A1-00', 'C1-00']})
    check('elenco diverso da quello confermato: non si manda niente (409)',
          r9.status_code == 409 and not lanciati, (r9.status_code, r9.get_json()))
    r9 = laser.post(f'/api/orders/{oid}/lantek-invia', json={'codici': ['A1-00']})
    j = r9.get_json() or {}
    check('confermato: XML Importer lanciato una volta con l\'XML di A1, verificato in Lantek',
          r9.status_code == 200 and len(lanciati) == 1 and b'A1-00' in lanciati[0][0]
          and b'B1-00' not in lanciati[0][0] and j.get('mandati') == 1 and j.get('verificati') == 1
          and j.get('mancanti') == [], (r9.status_code, j))
    r9 = laser.post(f'/api/orders/{oid}/lantek-invia', json={'codici': ['A1-00']})
    check('rimandato: ora e\' gia\' in produzione, niente da mandare (409), niente doppioni',
          r9.status_code == 409 and len(lanciati) == 1, (r9.status_code, r9.get_json()))
    check('l\'officina non puo\' mandare', rep.post(f'/api/orders/{oid}/lantek-invia',
                                                   json={'codici': ['A1-00']}).status_code in (401, 403))

    print('\n10) Pezzi nuovi col disegno (Procesos.exe, come iErp)')
    # formato provato in Lantek il 07/10/2026 (FT-PROVA-23/24)
    riga = L.riga_lista_dxf('FT-PROVA-23', r'C:\x\FT-PROVA-23.dxf', 'FERRO', 3.0, user_data=['PROVA FT', '15/10/2026'])
    check('riga dell\'elenco come quella provata',
          riga == '"FT-PROVA-23" 0 "" "" "" "CY Laser 3015 HL ECS" "FERRO" 3 1 "C:\\x\\FT-PROVA-23.dxf" '
                  '"PROVA FT" "15/10/2026" "" "" "" "" "" "" ', riga)
    check('spessore 1,5 col punto', ' 1.5 1 ' in L.riga_lista_dxf('X', 'x.dxf', 'INOX', 1.5))
    prc = L.processo_import_dxf(r'C:\x\DxfLista.Lst', 'FTIMP1')
    check('processo come quello provato', prc.splitlines() == [
        '0 FILEPROLT 8.02', '2', '1 1', '3 1 "FTIMP1" "FTIMP1" "" "" "" "" "" "" ""',
        '107 1 "C:\\x\\DxfLista.Lst"', '5', '38 1 "FTIMP1" 0 0', '2'], prc)
    # C1-00 e' nuovo: col suo DXF pronto si puo' creare (ALLUMINIO 4)
    pronto = os.path.join(_CARTELLE, 'C1-00.dxf')
    import ezdxf as _ez2
    _d3 = _ez2.new(); _d3.modelspace().add_lwpolyline([(0, 0), (90, 0), (90, 40), (0, 40)], close=True)
    _d3.saveas(pronto)
    A._struttura_zip = lambda radice, cartella, disegni, righe: [(pronto, 'X/ALU - 4 mm/C1-00.dxf')]
    A._disegni_ordine = lambda order: [{'nome': 'C1-00.dxf', 'percorso': pronto}]
    L.procesos_disponibile = lambda: True
    L.ordini_in_lantek = lambda commessa: {}
    pezzi_lantek = {'C1-00': {'esiste': False, 'codice_lantek': None, 'materiale': None, 'spessore': None, 'revisioni': []}}
    vecchio_pil = L.pezzi_in_lantek

    def pil(codici):
        e = vecchio_pil([c for c in codici if c != 'C1-00'])
        e = {**e, 'pezzi': dict(e.get('pezzi') or {})}
        if 'C1-00' in codici:
            e['pezzi']['C1-00'] = pezzi_lantek['C1-00']
        return e
    L.pezzi_in_lantek = pil
    d = laser.get(f'/api/orders/{oid}/lantek-quantita').get_json() or {}
    check('C1-00 nuovo: da creare col disegno, ALLUMINIO 4',
          [(x['codice'], x['materiale'], x['spessore']) for x in d.get('nuovi_da_creare') or []] == [('C1-00', 'ALLUMINIO', 4.0)],
          d.get('nuovi_da_creare'))
    passi = []

    def pezzi_finti(pezzi, cartella, user_data=None, attesa_s=900):
        passi.append(('pezzi', [p['codice'] for p in pezzi], [p['dxf'] for p in pezzi], user_data))
        pezzi_lantek['C1-00'] = {'esiste': True, 'codice_lantek': 'C1-00', 'materiale': 'ALLUMINIO',
                                 'spessore': 4.0, 'revisioni': []}
        return {'eseguito': True, 'errore': None, 'cartella': cartella}

    def xml_finto(contenuto, nome, cartella, attesa_s=300):
        passi.append(('xml', contenuto))
        L.ordini_in_lantek = lambda commessa: {'A1-00': 6.0, 'B1-00': 8.0, 'C1-00': 2.0}
        return {'eseguito': True, 'errore': None, 'file_rapporto': 'x',
                'rapporto': {'totale': 3, 'ok': 3, 'avvisi': 0, 'con_errore': 0, 'errori': []}}
    L.importa_pezzi_dxf = pezzi_finti
    L.importa_xml = xml_finto
    r10 = laser.post(f'/api/orders/{oid}/lantek-invia', json={'codici': sorted(x['codice'] for x in d['da_mandare'])})
    check('senza i nuovi visti in conferma: non si manda niente (409)', r10.status_code == 409 and not passi,
          (r10.status_code, r10.get_json()))
    r10 = laser.post(f'/api/orders/{oid}/lantek-invia',
                     json={'codici': [x['codice'] for x in d['da_mandare']], 'nuovi': ['C1-00']})
    j = r10.get_json() or {}
    check('prima i pezzi nuovi (C1-00 col suo DXF, cliente DECA e consegna nei dati), poi gli ordini',
          r10.status_code == 200 and [p[0] for p in passi] == ['pezzi', 'xml']
          and passi[0][1] == ['C1-00'] and passi[0][2] == [pronto] and passi[0][3] == ['DECA', '06/11/2030'],
          (r10.status_code, j, passi[:1]))
    check('negli ordini anche il pezzo appena creato (C1-00 x 2)',
          b'"C1-00"' in passi[1][1] and j.get('pezzi_creati') == 1 and j.get('verificati') == j.get('mandati') == 3,
          j)

    print('\n11) Verifica tecnica dei disegni (al caricamento dell\'ordine)')
    cp = A._controllo_pezzo
    check('abbinato per somiglianza e non confermato: da confermare',
          (cp({'dxf_filename': 'X-01.dxf', 'abbinamento': {'tipo': 'somiglianza', 'confermato': False}}) or {}).get('stato')
          == 'da_confermare')
    check('confermato o nome uguale: niente da fare',
          cp({'dxf_filename': 'X-01.dxf', 'abbinamento': {'tipo': 'somiglianza', 'confermato': True}}) is None
          and cp({'dxf_filename': 'X.dxf', 'abbinamento': {'tipo': 'esatto', 'confermato': True}}) is None)
    check('verifica automatica "da guardare": col motivo',
          cp({'esito_verifica': {'stato': 'da_guardare', 'motivi': ['area 40% diversa da Lantek']}})
          == {'stato': 'da_guardare', 'motivi': ['area 40% diversa da Lantek']})
    check('contorno corretto in automatico: da confermare',
          (cp({'contorno_auto': {'stato': 'da_confermare'}}) or {}).get('stato') == 'da_confermare')
    # ordine con C1-00 nuovo il cui disegno e' stato abbinato solo per somiglianza
    pr = PreventivoManager.create('DECA S.r.l.', 'paolo', quantita=1, numero_ordine_cliente='7777')
    pid2 = pr['id'] if isinstance(pr, dict) else pr
    PreventivoManager.replace_articoli(pid2, [
        {'codice': 'C1-00', 'quantita': 2, 'materiale': 'ALU', 'spessore_mm': 4, 'costo_base_override': 3,
         'dxf_filename': 'C1-00.dxf', 'abbinamento': {'tipo': 'somiglianza', 'file': 'C1-01.dxf', 'score': 0.8,
                                                       'confermato': False}}])
    s = models.SessionLocal()
    s.query(Preventivo).filter(Preventivo.id == pid2).first().status = 'INVIATO'
    s.commit(); s.close()
    oid2 = PreventivoManager.accetta_e_crea_ordine(pid2, 'paolo').get('order_id')
    L.ordini_in_lantek = lambda commessa: {}
    pezzi_lantek['C1-00'] = {'esiste': False, 'codice_lantek': None, 'materiale': None, 'spessore': None, 'revisioni': []}
    d = laser.get(f'/api/orders/{oid2}/lantek-quantita').get_json() or {}
    fuori = {x['codice']: x for x in d.get('nuovi_esclusi') or []}
    check('il pezzo nuovo dubbio non si crea: e\' fra quelli da guardare, col suo id',
          'C1-00' in fuori and fuori['C1-00'].get('verifica') and fuori['C1-00'].get('articolo_id')
          and not d.get('nuovi_da_creare'), d.get('nuovi_esclusi'))
    rc = laser.post(f"/api/orders/{oid2}/pezzi/{fuori.get('C1-00', {}).get('articolo_id', 0)}/conferma", json={})
    check('"il disegno e\' giusto": salvato', rc.status_code == 200, (rc.status_code, rc.get_json()))
    d = laser.get(f'/api/orders/{oid2}/lantek-quantita').get_json() or {}
    check('dopo la conferma si crea col suo disegno',
          [x['codice'] for x in d.get('nuovi_da_creare') or []] == ['C1-00'] and not d.get('nuovi_esclusi'),
          (d.get('nuovi_da_creare'), d.get('nuovi_esclusi')))
    rs = laser.get(f"/api/orders/{oid2}/pezzi/{fuori.get('C1-00', {}).get('articolo_id', 0)}/contorno.svg")
    corpo_svg = rs.data.decode('utf-8', 'replace')
    check('disegno col contorno preso evidenziato (verde)', rs.status_code == 200 and corpo_svg.startswith('<svg')
          and '#16a34a' in corpo_svg and '<polyline' in corpo_svg, (rs.status_code, corpo_svg[:160]))
    check('il tablet officina non vede il contorno',
          rep.get(f"/api/orders/{oid2}/pezzi/x/contorno.svg").status_code in (401, 403))
    check('l\'officina non puo\' confermare',
          rep.post(f"/api/orders/{oid2}/pezzi/1/conferma", json={}).status_code in (401, 403))

    print(f'\nPASSATI: {OK}   FALLITI: {len(KO)}')
    return 0 if not KO else 1


if __name__ == '__main__':
    try:
        sys.exit(main())
    finally:
        try:
            _ENG.dispose(); os.remove(_TMP)
            shutil.rmtree(_CARTELLE, ignore_errors=True)
        except Exception:
            pass
